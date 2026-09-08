from test_new_business_features import create_product, login, make_client
from test_workflows import create_unit_user_product_order


def test_authenticated_category_catalog_is_active_only_and_stably_sorted(tmp_path):
    client = make_client(tmp_path)
    admin_headers = login(client, "root_admin", "StrongPassword123")
    catalog = client.get("/api/v1/admin/product-categories", headers=admin_headers).json()
    frozen = next(item for item in catalog if item["name"] == "冻货")
    vegetables = next(item for item in catalog if item["name"] == "蔬菜")
    assert client.patch(f"/api/v1/admin/product-categories/{frozen['id']}", headers=admin_headers, json={"sort_order": 10}).status_code == 200
    assert client.patch(f"/api/v1/admin/product-categories/{vegetables['id']}", headers=admin_headers, json={"sort_order": 10}).status_code == 200
    unit_headers = create_unit_user_product_order(client)[1]

    assert client.get("/api/v1/product-categories").status_code == 401
    for headers in (admin_headers, unit_headers):
        response = client.get("/api/v1/product-categories", headers=headers)
        assert response.status_code == 200, response.text
        items = response.json()["items"]
        assert all(set(item) == {"id", "name", "sort_order"} for item in items)
        assert "冻货" in [item["name"] for item in items]
        assert [item["name"] for item in items].index("冻货") < [item["name"] for item in items].index("蔬菜")

    other = client.post("/api/v1/admin/product-categories", headers=admin_headers, json={"name": "临时停用分类", "sort_order": 1}).json()
    assert client.patch(f"/api/v1/admin/product-categories/{other['id']}", headers=admin_headers, json={"is_active": False}).status_code == 200
    assert "临时停用分类" not in [item["name"] for item in client.get("/api/v1/product-categories", headers=admin_headers).json()["items"]]


def test_category_catalog_crud_and_legacy_protection(tmp_path):
    client = make_client(tmp_path)
    headers = login(client, "root_admin", "StrongPassword123")

    listed = client.get("/api/v1/admin/product-categories", headers=headers)
    assert listed.status_code == 200
    names = [item["name"] for item in listed.json()]
    assert {"肉禽", "肉类", "蛋奶", "鸡蛋", "牛奶", "冻货"}.issubset(names)

    created = client.post("/api/v1/admin/product-categories", headers=headers, json={"name": " 水果干 ", "sort_order": 15})
    assert created.status_code == 200, created.text
    category = created.json()
    assert category["name"] == "水果干"
    assert client.post("/api/v1/admin/product-categories", headers=headers, json={"name": "水果干"}).status_code == 409

    renamed = client.patch(f"/api/v1/admin/product-categories/{category['id']}", headers=headers, json={"name": "干果", "sort_order": 16})
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["name"] == "干果"
    disabled = client.patch(f"/api/v1/admin/product-categories/{category['id']}", headers=headers, json={"is_active": False})
    assert disabled.status_code == 200
    assert disabled.json()["is_active"] is False
    assert client.patch(f"/api/v1/admin/product-categories/{category['id']}", headers=headers, json={"is_active": True}).status_code == 200

    protected = next(item for item in listed.json() if item["name"] == "肉禽")
    assert client.patch(f"/api/v1/admin/product-categories/{protected['id']}", headers=headers, json={"name": "肉类"}).status_code == 409
    assert client.patch(f"/api/v1/admin/product-categories/{protected['id']}", headers=headers, json={"is_active": False}).status_code == 409


def test_product_category_validation_preserves_inactive_current_value(tmp_path):
    client = make_client(tmp_path)
    headers = login(client, "root_admin", "StrongPassword123")
    category = client.post("/api/v1/admin/product-categories", headers=headers, json={"name": "临时分类"}).json()
    product = create_product(client, headers, code="DYNAMIC-CATEGORY", name="测试食材", unit="斤", price_cents=250, category="临时分类")
    assert client.patch(f"/api/v1/admin/product-categories/{category['id']}", headers=headers, json={"is_active": False}).status_code == 200
    preserved = client.put(f"/api/v1/admin/products/{product['id']}", headers=headers, json={"name": "测试食材二", "expected_version": product["version"]})
    assert preserved.status_code == 200, preserved.text
    assert preserved.json()["category"] == "临时分类"
    rejected = client.post("/api/v1/admin/products", headers=headers, json={
        "product_code": "DYNAMIC-CATEGORY-2", "name": "不能使用停用分类", "category": "临时分类",
        "spec": "散装", "unit": "斤", "price_cents": 250, "stock_quantity": "0",
        "min_order_quantity": "1", "quantity_step": "1", "supply_status": "paused", "active": True,
    })
    assert rejected.status_code == 400


def test_menu_export_uses_catalog_order_and_keeps_unknown_categories(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook

    client = make_client(tmp_path)
    headers = login(client, "root_admin", "StrongPassword123")
    catalog = client.get("/api/v1/admin/product-categories", headers=headers).json()
    frozen = next(item for item in catalog if item["name"] == "冻货")
    vegetables = next(item for item in catalog if item["name"] == "蔬菜")
    assert client.patch(f"/api/v1/admin/product-categories/{frozen['id']}", headers=headers, json={"sort_order": 5}).status_code == 200
    assert client.patch(f"/api/v1/admin/product-categories/{vegetables['id']}", headers=headers, json={"sort_order": 10}).status_code == 200
    create_product(client, headers, code="FROZEN-SHRIMP", name="冻虾", unit="斤", price_cents=1000, category="冻货")
    create_product(client, headers, code="VEG-CABBAGE", name="白菜", unit="斤", price_cents=200, category="蔬菜")
    exported = client.get("/api/v1/admin/products/export.xlsx", headers=headers)
    sheet = load_workbook(BytesIO(exported.content), data_only=True).active
    assert [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["冻货", "蔬菜"]
